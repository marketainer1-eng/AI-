"""
DocumentRenderer — produces a text-level preview of a DOCX document
with placeholder tokens replaced by rendered evidence numbers.

Key rule: this renderer does NOT modify the source DOCX file.
It produces:
  - a list of paragraph strings (for preview)
  - an output DOCX (for export) with replacements applied to a copy

File rename on disk happens ONLY in the commit flow, not here.
"""

import re
import shutil
from pathlib import Path
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from copy import deepcopy


class DocumentRenderer:
    """
    Renders a DOCX document by substituting placeholder tokens with
    their resolved evidence number strings.

    placeholder_map: {placeholder_text -> rendered_number_string}
    e.g. {"{{갑 제1호증}}": "갑 제1호증", "{{을 제2호증}}": "을 제2호증"}
    """

    def render_to_paragraphs(
        self,
        docx_path: Path,
        placeholder_map: dict[str, str],
    ) -> list[str]:
        """
        Return paragraph texts with placeholders substituted.
        This is the preview — the source file is not modified.
        """
        doc = DocxDocument(str(docx_path))
        result: list[str] = []

        for paragraph in doc.paragraphs:
            text = paragraph.text
            for placeholder, rendered in placeholder_map.items():
                text = text.replace(placeholder, rendered)
            result.append(text)

        return result

    def export_docx(
        self,
        docx_path: Path,
        placeholder_map: dict[str, str],
        output_path: Path,
        paragraphs: list[str] | None = None,
    ) -> Path:
        """
        Write a new DOCX file to output_path with placeholders replaced.
        Source file at docx_path is NOT modified.
        """
        # Copy source DOCX to output path
        shutil.copy2(str(docx_path), str(output_path))

        doc = DocxDocument(str(output_path))

        for paragraph in doc.paragraphs:
            if not placeholder_map:
                break
            # Replace in each run to preserve formatting
            for run in paragraph.runs:
                for placeholder, rendered in placeholder_map.items():
                    if placeholder in run.text:
                        run.text = run.text.replace(placeholder, rendered)

        doc.save(str(output_path))
        return output_path
