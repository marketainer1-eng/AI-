"""
EvidenceListRenderer — produces an evidence list preview and DOCX export.

Number computation rule:
  - Evidences are sorted by sort_order (ascending) within each party.
  - Rank (1-based) within the sorted party list = rendered number.
  - The rendered number is a PROJECTION — not stored on the Evidence entity.
"""

from pathlib import Path
from docx import Document as DocxDocument
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.api.schemas.preview import EvidenceListEntryPreview


class EvidenceListRenderer:
    def render(
        self,
        case_id: int,
        evidence_repo,
        number_map: dict[int, str],
    ) -> list[EvidenceListEntryPreview]:
        """
        Build evidence list preview entries.
        number_map: {evidence_id -> rendered_number_string}
        """
        entries: list[EvidenceListEntryPreview] = []
        evidences = evidence_repo.get_by_case(case_id)

        for evidence in sorted(evidences, key=lambda e: (e.party, e.sort_order)):
            rendered_number = number_map.get(evidence.id, "?")
            entries.append(
                EvidenceListEntryPreview(
                    evidence_id=evidence.id,
                    rendered_number=rendered_number,
                    rendered_label=evidence.label,
                    rendered_description=evidence.description,
                    sort_order_snapshot=evidence.sort_order,
                )
            )

        return entries

    def export_docx(
        self,
        entries: list[EvidenceListEntryPreview],
        output_path: Path,
    ) -> Path:
        """
        Write a DOCX evidence list to output_path.
        Format: [번호] [제목] [설명]
        """
        doc = DocxDocument()
        doc.add_heading("증거 목록", level=1)

        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "번호"
        hdr[1].text = "제목"
        hdr[2].text = "설명"

        for entry in entries:
            row = table.add_row().cells
            row[0].text = entry.rendered_number
            row[1].text = entry.rendered_label
            row[2].text = entry.rendered_description or ""

        doc.save(str(output_path))
        return output_path
