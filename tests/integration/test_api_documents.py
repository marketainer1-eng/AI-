"""
Integration tests for document-related API endpoints.

Endpoints covered
-----------------
POST /cases/{case_id}/documents                              – register
GET  /cases/{case_id}/documents                              – list
GET  /cases/{case_id}/documents/{doc_id}                    – get
POST /cases/{case_id}/documents/{doc_id}/parse-placeholders – parse
GET  /cases/{case_id}/documents/{doc_id}/anchors            – list anchors

Scenarios
---------
Happy path
  - register → 201 with correct fields
  - list returns items + total
  - get returns the registered document
  - parse on a real DOCX → 200 with anchors_found > 0
  - anchors endpoint returns parsed anchors

Edge / failure
  - register with missing case → 404
  - register with missing source_file → 404
  - get non-existent doc → 404
  - parse on non-existent document → 404
"""

from __future__ import annotations

import io
import tempfile
from pathlib import Path
from fastapi.testclient import TestClient


# ── helpers ───────────────────────────────────────────────────────────────────

def _create_case(client: TestClient, name: str = "문서 테스트 사건") -> int:
    resp = client.post("/cases", json={"name": name})
    assert resp.status_code == 201
    return resp.json()["id"]


def _upload_file(
    client: TestClient,
    case_id: int,
    content: bytes = b"dummy",
    filename: str = "brief.docx",
    role: str = "document",
) -> int:
    resp = client.post(
        f"/cases/{case_id}/files",
        params={"role": role},
        files={"file": (filename, io.BytesIO(content), "application/octet-stream")},
    )
    assert resp.status_code == 201
    return resp.json()["id"]


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    """Create an in-memory DOCX and return its bytes."""
    from docx import Document as DocxDoc
    doc = DocxDoc()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════

class TestDocumentRegister:
    def test_register_returns_201(self, client: TestClient):
        case_id = _create_case(client)
        file_id = _upload_file(client, case_id)
        resp = client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": file_id, "title": "준비서면", "doc_type": "main_brief"},
        )
        assert resp.status_code == 201

    def test_register_response_fields(self, client: TestClient):
        case_id = _create_case(client)
        file_id = _upload_file(client, case_id)
        resp = client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": file_id, "title": "2024 준비서면", "doc_type": "main_brief"},
        )
        data = resp.json()
        assert data["id"] is not None
        assert data["title"] == "2024 준비서면"
        assert data["doc_type"] == "main_brief"
        assert data["parse_status"] == "pending"

    def test_register_case_not_found(self, client: TestClient):
        resp = client.post(
            "/cases/99999/documents",
            json={"source_file_id": 1, "title": "X", "doc_type": "main_brief"},
        )
        assert resp.status_code == 404

    def test_register_source_file_not_found(self, client: TestClient):
        case_id = _create_case(client)
        resp = client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": 99999, "title": "X", "doc_type": "main_brief"},
        )
        assert resp.status_code == 404

    def test_register_invalid_doc_type(self, client: TestClient):
        case_id = _create_case(client)
        file_id = _upload_file(client, case_id)
        resp = client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": file_id, "title": "X", "doc_type": "invalid_type"},
        )
        assert resp.status_code == 422


class TestDocumentList:
    def test_list_empty(self, client: TestClient):
        case_id = _create_case(client)
        resp = client.get(f"/cases/{case_id}/documents")
        assert resp.status_code == 200
        data = resp.json()
        assert data["total"] == 0
        assert data["items"] == []

    def test_list_after_register(self, client: TestClient):
        # documents.source_file_id has a UNIQUE constraint — use separate files
        case_id = _create_case(client)
        file_id_a = _upload_file(
            client, case_id,
            content=b"unique-content-for-doc-A",
            filename="brief_a.docx",
        )
        file_id_b = _upload_file(
            client, case_id,
            content=b"unique-content-for-doc-B",
            filename="exhibit_b.docx",
        )
        client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": file_id_a, "title": "A", "doc_type": "main_brief"},
        )
        client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": file_id_b, "title": "B", "doc_type": "exhibit_list"},
        )
        resp = client.get(f"/cases/{case_id}/documents")
        assert resp.json()["total"] == 2

    def test_list_case_not_found(self, client: TestClient):
        resp = client.get("/cases/99999/documents")
        assert resp.status_code == 404


class TestDocumentGet:
    def test_get_document(self, client: TestClient):
        case_id = _create_case(client)
        file_id = _upload_file(client, case_id)
        doc_resp = client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": file_id, "title": "T", "doc_type": "main_brief"},
        )
        doc_id = doc_resp.json()["id"]
        resp = client.get(f"/cases/{case_id}/documents/{doc_id}")
        assert resp.status_code == 200
        assert resp.json()["id"] == doc_id

    def test_get_document_not_found(self, client: TestClient):
        case_id = _create_case(client)
        resp = client.get(f"/cases/{case_id}/documents/99999")
        assert resp.status_code == 404


class TestParsePlaceholders:
    def test_parse_with_placeholders(self, client: TestClient):
        """Upload a real DOCX with placeholders and verify anchors are extracted."""
        case_id = _create_case(client)
        docx_bytes = _make_docx_bytes([
            "{{갑 제1호증}} 참조.",
            "{{갑 제2호증}} 참조.",
        ])
        file_id = _upload_file(
            client, case_id, content=docx_bytes, filename="brief.docx"
        )
        doc_resp = client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": file_id, "title": "준비서면", "doc_type": "main_brief"},
        )
        doc_id = doc_resp.json()["id"]

        resp = client.post(f"/cases/{case_id}/documents/{doc_id}/parse-placeholders")
        assert resp.status_code == 200
        data = resp.json()
        assert data["anchors_found"] == 2
        assert len(data["anchors"]) == 2

    def test_parse_updates_parse_status(self, client: TestClient):
        case_id = _create_case(client)
        docx_bytes = _make_docx_bytes(["{{갑 제1호증}}"])
        file_id = _upload_file(client, case_id, content=docx_bytes)
        doc_resp = client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": file_id, "title": "T", "doc_type": "main_brief"},
        )
        doc_id = doc_resp.json()["id"]
        client.post(f"/cases/{case_id}/documents/{doc_id}/parse-placeholders")

        doc_detail = client.get(f"/cases/{case_id}/documents/{doc_id}").json()
        assert doc_detail["parse_status"] == "parsed"

    def test_parse_no_placeholders(self, client: TestClient):
        case_id = _create_case(client)
        docx_bytes = _make_docx_bytes(["플레이스홀더 없는 문단"])
        file_id = _upload_file(client, case_id, content=docx_bytes)
        doc_resp = client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": file_id, "title": "T", "doc_type": "main_brief"},
        )
        doc_id = doc_resp.json()["id"]
        resp = client.post(f"/cases/{case_id}/documents/{doc_id}/parse-placeholders")
        assert resp.status_code == 200
        assert resp.json()["anchors_found"] == 0

    def test_parse_doc_not_found(self, client: TestClient):
        case_id = _create_case(client)
        resp = client.post(f"/cases/{case_id}/documents/99999/parse-placeholders")
        assert resp.status_code == 404

    def test_parse_idempotent(self, client: TestClient):
        """Calling parse twice should replace anchors, not duplicate them."""
        case_id = _create_case(client)
        docx_bytes = _make_docx_bytes(["{{갑 제1호증}}", "{{갑 제2호증}}"])
        file_id = _upload_file(client, case_id, content=docx_bytes)
        doc_resp = client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": file_id, "title": "T", "doc_type": "main_brief"},
        )
        doc_id = doc_resp.json()["id"]
        client.post(f"/cases/{case_id}/documents/{doc_id}/parse-placeholders")
        resp2 = client.post(f"/cases/{case_id}/documents/{doc_id}/parse-placeholders")
        assert resp2.json()["anchors_found"] == 2


class TestAnchorsList:
    def test_anchors_empty_before_parse(self, client: TestClient):
        case_id = _create_case(client)
        file_id = _upload_file(client, case_id)
        doc_resp = client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": file_id, "title": "T", "doc_type": "main_brief"},
        )
        doc_id = doc_resp.json()["id"]
        resp = client.get(f"/cases/{case_id}/documents/{doc_id}/anchors")
        assert resp.status_code == 200
        assert resp.json() == []

    def test_anchors_after_parse(self, client: TestClient):
        case_id = _create_case(client)
        docx_bytes = _make_docx_bytes(["{{갑 제1호증}}"])
        file_id = _upload_file(client, case_id, content=docx_bytes)
        doc_resp = client.post(
            f"/cases/{case_id}/documents",
            json={"source_file_id": file_id, "title": "T", "doc_type": "main_brief"},
        )
        doc_id = doc_resp.json()["id"]
        client.post(f"/cases/{case_id}/documents/{doc_id}/parse-placeholders")
        resp = client.get(f"/cases/{case_id}/documents/{doc_id}/anchors")
        assert resp.status_code == 200
        anchors = resp.json()
        assert len(anchors) == 1
        assert anchors[0]["placeholder_text"] == "{{갑 제1호증}}"
        assert anchors[0]["status"] == "unlinked"
