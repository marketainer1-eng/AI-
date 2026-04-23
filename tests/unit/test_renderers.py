"""
Unit tests for renderers.
"""

import tempfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.renderers.document_renderer import DocumentRenderer
from app.renderers.filename_renderer import FilenameRenderer, sanitize_filename
from app.renderers.evidence_list_renderer import EvidenceListRenderer
from app.api.schemas.preview import EvidenceListEntryPreview


def _make_docx(paragraphs: list[str]) -> Path:
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    doc.save(tmp.name)
    tmp.close()
    return Path(tmp.name)


class TestDocumentRenderer:
    def test_render_to_paragraphs_substitutes_placeholder(self):
        path = _make_docx(["이 사건은 {{갑 제1호증}} 참조."])
        renderer = DocumentRenderer()
        result = renderer.render_to_paragraphs(
            docx_path=path,
            placeholder_map={"{{갑 제1호증}}": "갑 제1호증"},
        )
        assert any("갑 제1호증" in p for p in result)
        assert not any("{{갑 제1호증}}" in p for p in result)

    def test_render_to_paragraphs_multiple_substitutions(self):
        path = _make_docx([
            "{{갑 제1호증}} 및",
            "{{갑 제2호증}} 참조.",
        ])
        renderer = DocumentRenderer()
        result = renderer.render_to_paragraphs(
            docx_path=path,
            placeholder_map={
                "{{갑 제1호증}}": "갑 제1호증",
                "{{갑 제2호증}}": "갑 제2호증",
            },
        )
        assert any("갑 제1호증" in p for p in result)
        assert any("갑 제2호증" in p for p in result)

    def test_render_to_paragraphs_no_placeholder_map(self):
        path = _make_docx(["일반 문단", "{{갑 제1호증}}"])
        renderer = DocumentRenderer()
        result = renderer.render_to_paragraphs(docx_path=path, placeholder_map={})
        assert "{{갑 제1호증}}" in result

    def test_export_docx_creates_file(self, tmp_path: Path):
        src_path = _make_docx(["{{갑 제1호증}} 참조."])
        out_path = tmp_path / "output.docx"
        renderer = DocumentRenderer()
        renderer.export_docx(
            docx_path=src_path,
            placeholder_map={"{{갑 제1호증}}": "갑 제1호증"},
            output_path=out_path,
        )
        assert out_path.exists()

    def test_export_docx_does_not_modify_source(self, tmp_path: Path):
        src_path = _make_docx(["{{갑 제1호증}}"])
        out_path = tmp_path / "output.docx"
        renderer = DocumentRenderer()
        renderer.export_docx(
            docx_path=src_path,
            placeholder_map={"{{갑 제1호증}}": "갑 제1호증"},
            output_path=out_path,
        )
        # Source still has placeholder
        doc = DocxDocument(str(src_path))
        texts = [p.text for p in doc.paragraphs]
        assert any("{{갑 제1호증}}" in t for t in texts)


class TestSanitizeFilename:
    def test_removes_illegal_chars(self):
        assert sanitize_filename('계약서:사본<1>.docx') == '계약서_사본_1_.docx'

    def test_clean_name_unchanged(self):
        assert sanitize_filename("갑 제1호증_계약서.docx") == "갑 제1호증_계약서.docx"


class TestEvidenceListRenderer:
    def test_export_docx_creates_file(self, tmp_path: Path):
        entries = [
            EvidenceListEntryPreview(
                evidence_id=1,
                rendered_number="갑 제1호증",
                rendered_label="계약서",
                rendered_description="사본",
                sort_order_snapshot=1,
            ),
            EvidenceListEntryPreview(
                evidence_id=2,
                rendered_number="갑 제2호증",
                rendered_label="영수증",
                rendered_description=None,
                sort_order_snapshot=2,
            ),
        ]
        out_path = tmp_path / "evidence_list.docx"
        renderer = EvidenceListRenderer()
        renderer.export_docx(entries=entries, output_path=out_path)
        assert out_path.exists()

    def test_export_docx_contains_entries(self, tmp_path: Path):
        entries = [
            EvidenceListEntryPreview(
                evidence_id=1,
                rendered_number="갑 제1호증",
                rendered_label="계약서",
                rendered_description=None,
                sort_order_snapshot=1,
            ),
        ]
        out_path = tmp_path / "ev.docx"
        renderer = EvidenceListRenderer()
        renderer.export_docx(entries=entries, output_path=out_path)
        doc = DocxDocument(str(out_path))
        full_text = " ".join(p.text for p in doc.paragraphs)
        # Table content may not appear in paragraphs — check tables
        table_text = ""
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    table_text += cell.text + " "
        assert "갑 제1호증" in table_text
        assert "계약서" in table_text
